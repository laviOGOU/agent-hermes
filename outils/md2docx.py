"""
md2docx.py — convertit un fichier Markdown en document Word (.docx).

Usage :
    uv run --with python-docx python md2docx.py fichier.md [autre.md ...]
    uv run --with python-docx python md2docx.py --dossier D:\\chemin\\dossier

Gère : titres (#, ##, ###), gras (**texte**), code (`texte`), listes à puces,
listes numérotées, citations (>), tableaux (| a | b |), blocs de code (```),
et met en évidence les champs à compléter de la forme [EN MAJUSCULES].

Le fichier .docx est écrit à côté du .md, avec le même nom.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

POLICE = "Calibri"
POLICE_TITRE = "Georgia"
POLICE_CODE = "Consolas"
ROUGE_CHAMP = RGBColor(0xC0, 0x39, 0x2B)
GRIS_CITATION = RGBColor(0x55, 0x55, 0x55)

CHAMP = re.compile(r"\[[A-ZÉÈÀÊÎÔÛÇ][A-ZÉÈÀÊÎÔÛÇ0-9 \-'’/]{2,}\]")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[A-ZÉÈÀÊÎÔÛÇ][^\]]{2,}\])")


def preparer_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = POLICE
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
    for nom, taille in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        st = doc.styles[nom]
        st.font.name = POLICE_TITRE
        st.font.size = Pt(taille)
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def ajouter_runs(paragraphe, texte):
    """Écrit le texte en interprétant **gras**, `code` et [CHAMP À COMPLÉTER]."""
    for morceau in INLINE.split(texte):
        if not morceau:
            continue
        if morceau.startswith("**") and morceau.endswith("**") and len(morceau) > 4:
            paragraphe.add_run(morceau[2:-2]).bold = True
        elif morceau.startswith("`") and morceau.endswith("`") and len(morceau) > 2:
            run = paragraphe.add_run(morceau[1:-1])
            run.font.name = POLICE_CODE
            run.font.size = Pt(9.5)
        elif CHAMP.fullmatch(morceau):
            run = paragraphe.add_run(morceau)
            run.bold = True
            run.font.color.rgb = ROUGE_CHAMP
        else:
            paragraphe.add_run(morceau)


def tableau(doc, lignes):
    cellules = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lignes]
    cellules = [l for l in cellules if not all(set(c) <= set("-: ") for c in l)]
    if not cellules:
        return
    largeur = max(len(l) for l in cellules)
    t = doc.add_table(rows=0, cols=largeur)
    t.style = "Table Grid"
    for i, ligne in enumerate(cellules):
        cellule = t.add_row().cells
        for j in range(largeur):
            valeur = ligne[j] if j < len(ligne) else ""
            p = cellule[j].paragraphs[0]
            ajouter_runs(p, valeur)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def convertir(md: Path, docx: Path) -> None:
    doc = Document()
    preparer_styles(doc)
    lignes = md.read_text(encoding="utf-8").splitlines()

    dans_code = False
    tampon_tableau = []
    tampon_code = []

    def vider_tableau():
        nonlocal tampon_tableau
        if tampon_tableau:
            tableau(doc, tampon_tableau)
            tampon_tableau = []

    def vider_code():
        nonlocal tampon_code
        if tampon_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run("\n".join(tampon_code))
            run.font.name = POLICE_CODE
            run.font.size = Pt(9)
            tampon_code = []

    for brute in lignes:
        ligne = brute.rstrip()

        if ligne.strip().startswith("```"):
            if dans_code:
                vider_code()
                dans_code = False
            else:
                vider_tableau()
                dans_code = True
            continue
        if dans_code:
            tampon_code.append(brute)
            continue

        if ligne.strip().startswith("|"):
            tampon_tableau.append(ligne)
            continue
        vider_tableau()

        if not ligne.strip():
            continue
        if set(ligne.strip()) <= set("-—_ ") and len(ligne.strip()) >= 3:
            continue                                   # trait de séparation

        if ligne.startswith("#"):
            niveau = len(ligne) - len(ligne.lstrip("#"))
            texte = ligne[niveau:].strip()
            doc.add_heading(texte, level=min(niveau, 3))
            continue

        if ligne.lstrip().startswith(("- ", "• ")):
            ajouter_runs(doc.add_paragraph(style="List Bullet"), ligne.lstrip()[2:])
            continue
        if re.match(r"^\s*\d+\.\s", ligne):
            texte = re.sub(r"^\s*\d+\.\s", "", ligne)
            ajouter_runs(doc.add_paragraph(style="List Number"), texte)
            continue
        if ligne.lstrip().startswith("- [ ]"):
            ajouter_runs(doc.add_paragraph(style="List Bullet"), "☐ " + ligne.lstrip()[5:])
            continue

        if ligne.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(20)
            run_texte = ligne.lstrip()[1:].strip()
            ajouter_runs(p, run_texte)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = GRIS_CITATION
            continue

        ajouter_runs(doc.add_paragraph(), ligne)

    vider_tableau()
    vider_code()

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(56)
        section.left_margin = section.right_margin = Pt(60)

    doc.save(docx)
    print(f"  {md.name}  ->  {docx.name}")


def main(argv):
    fichiers = []
    if "--dossier" in argv:
        dossier = Path(argv[argv.index("--dossier") + 1])
        fichiers = sorted(dossier.glob("*.md"))
    else:
        fichiers = [Path(a) for a in argv[1:] if a.endswith(".md")]

    if not fichiers:
        print(__doc__)
        return 1

    for md in fichiers:
        try:
            convertir(md, md.with_suffix(".docx"))
        except Exception as exc:                       # noqa: BLE001
            print(f"  ECHEC {md.name} : {type(exc).__name__}: {exc}")
            return 1
    print(f"\n{len(fichiers)} document(s) converti(s).")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
